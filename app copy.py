
import io
import os
import json
import csv
import gc
import tempfile
import sqlite3
import warnings
from datetime import datetime

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, LongTable
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from sklearn.compose import ColumnTransformer
from sklearn.ensemble import (
    RandomForestClassifier, RandomForestRegressor,
    IsolationForest, ExtraTreesClassifier, ExtraTreesRegressor,
    HistGradientBoostingClassifier, HistGradientBoostingRegressor
)
from sklearn.impute import SimpleImputer
from sklearn.metrics import (
    accuracy_score, precision_score, recall_score, f1_score,
    confusion_matrix, classification_report,
    mean_absolute_error, mean_squared_error, r2_score,
    silhouette_score
)
from sklearn.model_selection import train_test_split, cross_val_score
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder, StandardScaler, LabelEncoder
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA
from sklearn.feature_selection import mutual_info_classif, mutual_info_regression
from sklearn.exceptions import ConvergenceWarning

warnings.filterwarnings("ignore", category=ConvergenceWarning)
st.set_page_config(page_title="AI Data Intelligence Pro", page_icon="🧠", layout="wide", initial_sidebar_state="expanded")

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Cairo:wght@400;500;600;700;800&display=swap');
html, body, [class*="css"] {font-family:'Cairo',sans-serif;}
.stApp {background:linear-gradient(135deg,#0b1220 0%,#101a2d 52%,#142338 100%);color:#e8eef8;}
[data-testid="stSidebar"] {background:linear-gradient(180deg,#0a1020,#111d31);border-right:1px solid #263852;}
[data-testid="stSidebar"] * {color:#e8eef8 !important;}
.block-container {padding-top:1.5rem; padding-bottom:2rem; max-width:1500px;}
h1,h2,h3 {color:#f8fafc !important; letter-spacing:-.3px;}
p, label, .stMarkdown {color:#d5dfed;}
.hero {background:linear-gradient(115deg,rgba(37,99,235,.25),rgba(20,184,166,.15));border:1px solid rgba(125,170,230,.3);border-radius:22px;padding:26px 28px;margin-bottom:18px;box-shadow:0 14px 35px rgba(0,0,0,.15);}
.hero-title {font-size:32px;font-weight:800;color:#fff;margin:0 0 4px;}
.hero-sub {font-size:15px;color:#b7c7dc;margin:0;}
.section-card {background:rgba(17,29,49,.82);border:1px solid #263852;border-radius:18px;padding:18px 20px;margin:10px 0;}
.metric-card {background:linear-gradient(145deg,#16253b,#101a2b);border:1px solid #2c405e;border-radius:17px;padding:17px 18px;min-height:115px;box-shadow:0 8px 20px rgba(0,0,0,.12);}
.metric-label {font-size:13px;color:#9fb2cc;margin-bottom:7px;}
.metric-value {font-size:28px;font-weight:800;color:#fff;line-height:1.25;}
.metric-note {font-size:12px;color:#62d9c3;margin-top:5px;}
.badge {display:inline-block;background:#123c42;color:#6ee7d3;border:1px solid #236b6c;border-radius:999px;padding:4px 11px;font-size:12px;font-weight:700;}
div.stButton > button {border-radius:11px;border:1px solid #315579;background:linear-gradient(90deg,#2563eb,#0d9488);color:white;font-weight:700;min-height:42px;}
div.stButton > button:hover {border-color:#72e3d0;box-shadow:0 0 0 2px rgba(45,212,191,.13);}
.stDownloadButton > button {border-radius:10px;background:#152a40;border:1px solid #35516f;color:#f8fafc;font-weight:600;}
[data-testid="stDataFrame"] {border:1px solid #293d57;border-radius:12px;overflow:hidden;}
div[data-baseweb="select"] > div, div[data-baseweb="input"] > div, .stTextInput input, .stNumberInput input {background:#101d30;color:#fff;border-color:#30445f;}
hr {border-color:#263852;}
.small-muted {font-size:12px;color:#8fa4bf;}
</style>
"""
st.markdown(CSS, unsafe_allow_html=True)

def init_state():
    defaults = {
        "df": None, "source_name": "", "clean_df": None,
        "trained_model": None, "model_info": None, "automl_results": None,
        "outlier_df": None, "cluster_df": None
    }
    for k, v in defaults.items():
        if k not in st.session_state:
            st.session_state[k] = v
init_state()

def metric_card(label, value, note=""):
    st.markdown(
        f'<div class="metric-card"><div class="metric-label">{label}</div>'
        f'<div class="metric-value">{value}</div><div class="metric-note">{note}</div></div>',
        unsafe_allow_html=True
    )

def read_dataset(uploaded):
    name = uploaded.name.lower()
    uploaded.seek(0)

    if name.endswith(".csv"):
        try:
            return pd.read_csv(uploaded, low_memory=True, on_bad_lines="error")
        except UnicodeDecodeError:
            uploaded.seek(0)
            return pd.read_csv(uploaded, encoding="latin-1", low_memory=True, on_bad_lines="error")

    if name.endswith((".xlsx", ".xlsm")):
        uploaded.seek(0)
        try:
            return pd.read_excel(uploaded, engine="calamine")
        except Exception as calamine_error:
            uploaded.seek(0)
            try:
                return pd.read_excel(uploaded, engine="openpyxl")
            except Exception as openpyxl_error:
                raise ValueError(
                    "تعذر قراءة ملف Excel. جرّب تثبيت python-calamine وopenpyxl، "
                    f"أو احفظ الملف بصيغة CSV. تفاصيل المحاولة السريعة: {calamine_error}; "
                    f"تفاصيل المحاولة البديلة: {openpyxl_error}"
                ) from openpyxl_error

    if name.endswith(".xls"):
        uploaded.seek(0)
        try:
            return pd.read_excel(uploaded, engine="calamine")
        except Exception:
            uploaded.seek(0)
            return pd.read_excel(uploaded, engine="xlrd")

    if name.endswith(".json"):
        uploaded.seek(0)
        return pd.read_json(uploaded)

    if name.endswith(".parquet"):
        uploaded.seek(0)
        return pd.read_parquet(uploaded)

    raise ValueError("الصيغة غير مدعومة. استخدم CSV أو XLSX أو XLS أو JSON أو Parquet.")


def safe_duplicate_count(frame, max_rows=200_000):
    """Avoid a large temporary hash array on very large datasets."""
    if len(frame) > max_rows:
        return None
    return int(frame.duplicated().sum())


def safe_unique_count(series, max_rows=200_000):
    if len(series) <= max_rows:
        return int(series.nunique(dropna=True)), False
    sample = series.sample(n=min(100_000, len(series)), random_state=42)
    return int(sample.nunique(dropna=True)), True


def plot_sample(frame, max_rows=50_000):
    if len(frame) <= max_rows:
        return frame
    return frame.sample(n=max_rows, random_state=42)

def classify_task(series):
    non_null = series.dropna()
    if non_null.empty:
        return "غير صالح"
    if pd.api.types.is_numeric_dtype(series):
        unique = non_null.nunique()
        if unique <= 12 and unique / max(len(non_null), 1) < 0.08:
            return "Classification"
        return "Regression"
    return "Classification"

def infer_target_candidates(df):
    candidates = []
    for col in df.columns:
        s = df[col]
        if s.notna().sum() >= 10 and s.nunique(dropna=True) >= 2:
            candidates.append(col)
    return candidates

def make_preprocessor(X):
    num_cols = X.select_dtypes(include=np.number).columns.tolist()
    cat_cols = [c for c in X.columns if c not in num_cols]
    transformers = []
    if num_cols:
        transformers.append(("num", Pipeline([
            ("imputer", SimpleImputer(strategy="median")),
            ("scaler", StandardScaler())
        ]), num_cols))
    if cat_cols:
        try:
            encoder = OneHotEncoder(handle_unknown="ignore", sparse_output=False)
        except TypeError:
            encoder = OneHotEncoder(handle_unknown="ignore", sparse=False)
        transformers.append(("cat", Pipeline([
            ("imputer", SimpleImputer(strategy="most_frequent")),
            ("encoder", encoder)
        ]), cat_cols))
    return ColumnTransformer(transformers=transformers, remainder="drop")

def safe_frame(df):
    datetime_cols = [c for c in df.columns if pd.api.types.is_datetime64_any_dtype(df[c])]
    if not datetime_cols:
        return df
    out = df.copy()
    for col in datetime_cols:
        out[col] = out[col].astype(str)
    return out

def to_csv_bytes(df):
    return safe_frame(df).to_csv(index=False).encode("utf-8-sig")

def to_excel_bytes(df):
    """Export original data plus a polished, dataset-aware Excel dashboard."""
    from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
    from openpyxl.chart import BarChart, LineChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.formatting.rule import DataBarRule
    from openpyxl.utils import get_column_letter

    original = df.copy() if isinstance(df, pd.DataFrame) else pd.DataFrame(df)
    if original.columns.duplicated().any():
        seen, names = {}, []
        for col in original.columns:
            key = str(col)
            seen[key] = seen.get(key, 0) + 1
            names.append(f"{key}_{seen[key]}" if seen[key] > 1 else key)
        original.columns = names

    work = original.copy()
    numeric_cols, date_cols = [], {}
    for col in work.columns:
        s = work[col]
        name = str(col).lower()
        is_date_name = any(x in name for x in ("date", "time", "التاريخ", "تاريخ", "وقت"))
        if pd.api.types.is_datetime64_any_dtype(s):
            parsed = pd.to_datetime(s, errors="coerce", format="mixed")
            if parsed.notna().mean() >= .6:
                date_cols[col] = parsed
        elif is_date_name and not pd.api.types.is_numeric_dtype(s):
            parsed = pd.to_datetime(s, errors="coerce", format="mixed")
            if len(s) and parsed.notna().mean() >= .6:
                date_cols[col] = parsed
        elif not pd.api.types.is_numeric_dtype(s) and s.notna().any():
            sample = s.dropna().astype(str).head(200)
            if len(sample) >= 5 and sample.str.contains(r"[-/]", regex=True).mean() >= .7:
                parsed = pd.to_datetime(s, errors="coerce", format="mixed")
                if parsed.notna().mean() >= .8:
                    date_cols[col] = parsed

    for col in work.columns:
        if col in date_cols:
            continue
        if pd.api.types.is_numeric_dtype(work[col]):
            numeric_cols.append(col)
        elif work[col].notna().any():
            converted = pd.to_numeric(work[col].astype(str).str.replace(",", "", regex=False), errors="coerce")
            if converted.notna().mean() >= .85:
                work[col] = converted
                numeric_cols.append(col)

    cat_cols = [c for c in work.columns if c not in numeric_cols and c not in date_cols]
    id_tokens = ("id", "code", "serial", "رقم", "كود")
    useful_numeric = [c for c in numeric_cols if not any(t in str(c).lower() for t in id_tokens)] or numeric_cols[:]
    nrows, ncols = original.shape
    missing_by_col = original.isna().sum()
    missing_total = int(missing_by_col.sum())
    duplicates = int(original.duplicated().sum())
    missing_pct = missing_total / max(nrows * ncols, 1) * 100
    buf = io.BytesIO()
    navy, teal, blue, orange, green, pale = "14263D", "147D82", "3977C3", "E07A45", "3A9D78", "F1F6FA"
    border = Side(style="thin", color="D8E2EB")

    with pd.ExcelWriter(buf, engine="openpyxl", datetime_format="yyyy-mm-dd hh:mm:ss", date_format="yyyy-mm-dd") as writer:
        safe_frame(original).to_excel(writer, index=False, sheet_name="Data")
        wb = writer.book
        dash = wb.create_sheet("Dashboard", 0)
        analysis = wb.create_sheet("Analysis")
        data_ws = writer.sheets["Data"]
        dash.sheet_view.showGridLines = False
        analysis.sheet_view.showGridLines = False
        data_ws.sheet_view.showGridLines = False

        # Summary / analysis tables
        analysis.append(["DATASET PROFILE", "VALUE"])
        for label, value in [
            ("Total records", nrows), ("Total columns", ncols), ("Missing cells", missing_total),
            ("Missing cells (%)", round(missing_pct, 2)), ("Duplicate rows", duplicates),
            ("Numeric columns", len(numeric_cols)), ("Categorical columns", len(cat_cols)),
            ("Date columns", len(date_cols))]:
            analysis.append([label, value])

        r = analysis.max_row + 3
        if numeric_cols:
            analysis.cell(r, 1, "NUMERIC STATISTICS")
            headers = ["Column", "Valid values", "Mean", "Median", "Minimum", "Maximum", "Std deviation"]
            for j, h in enumerate(headers, 1): analysis.cell(r + 1, j, h)
            for i, col in enumerate(numeric_cols, r + 2):
                s = pd.to_numeric(work[col], errors="coerce")
                vals = [str(col), int(s.notna().sum()), s.mean(), s.median(), s.min(), s.max(), s.std()]
                for j, val in enumerate(vals, 1):
                    if isinstance(val, np.generic): val = val.item()
                    if isinstance(val, float) and not np.isfinite(val): val = None
                    analysis.cell(i, j, val)
            r += len(numeric_cols) + 4

        if cat_cols:
            analysis.cell(r, 1, "CATEGORY BREAKDOWNS")
            r += 1
            for col in cat_cols:
                counts = work[col].fillna("(Missing)").astype(str).value_counts().head(10)
                analysis.cell(r, 1, f"Column: {col}")
                analysis.cell(r + 1, 1, str(col)); analysis.cell(r + 1, 2, "Records")
                for i, (label, count) in enumerate(counts.items(), r + 2):
                    analysis.cell(i, 1, str(label)[:100]); analysis.cell(i, 2, int(count))
                r += len(counts) + 4

        # Build sensible chart tables; never compare averages across unrelated units.
        charts = []
        src = max(r, analysis.max_row + 3)
        candidates = []
        for col in cat_cols:
            unique = int(work[col].nunique(dropna=True))
            if 2 <= unique <= min(25, max(2, nrows)) and not any(x in str(col).lower() for x in ("description", "comment", "address", "notes", "وصف", "عنوان")):
                candidates.append(col)
        if candidates:
            chosen = max(candidates, key=lambda c: min(work[c].nunique(dropna=True), 12))
            counts = work[chosen].fillna("(Missing)").astype(str).value_counts().head(10)
            counts = counts.sort_values(ascending=True)
            analysis.cell(src, 1, f"CHART DATA - TOP CATEGORIES: {chosen}")
            analysis.cell(src + 1, 1, str(chosen)); analysis.cell(src + 1, 2, "Record count")
            for i, (label, count) in enumerate(counts.items(), src + 2):
                analysis.cell(i, 1, str(label)[:80]); analysis.cell(i, 2, int(count))
            charts.append({"kind":"bar", "title":f"Records by {chosen}", "header":src+1, "count":len(counts), "color":teal,
                           "note":f"Compares the 10 most common values in '{chosen}'. Longer bars mean more records in that group."})
            src += len(counts) + 5

        miss = missing_by_col[missing_by_col > 0].sort_values(ascending=True).tail(10)
        if not miss.empty:
            analysis.cell(src, 1, "CHART DATA - MISSING VALUES")
            analysis.cell(src + 1, 1, "Column"); analysis.cell(src + 1, 2, "Blank cells")
            for i, (label, count) in enumerate(miss.items(), src + 2):
                analysis.cell(i, 1, str(label)[:80]); analysis.cell(i, 2, int(count))
            charts.append({"kind":"bar", "title":"Columns needing data cleanup", "header":src+1, "count":len(miss), "color":orange,
                           "note":"Shows the columns with the most empty cells. A longer bar means more missing values to investigate."})
            src += len(miss) + 5

        if date_cols:
            date_col = next(iter(date_cols))
            valid = date_cols[date_col].dropna()
            if len(valid):
                period_counts = valid.dt.to_period("M").astype(str).value_counts().sort_index().tail(24)
                analysis.cell(src, 1, f"CHART DATA - MONTHLY TREND: {date_col}")
                analysis.cell(src + 1, 1, "Month"); analysis.cell(src + 1, 2, "Records")
                for i, (label, count) in enumerate(period_counts.items(), src + 2):
                    analysis.cell(i, 1, str(label)); analysis.cell(i, 2, int(count))
                charts.append({"kind":"line", "title":f"Monthly record trend ({date_col})", "header":src+1, "count":len(period_counts), "color":blue,
                               "note":"Each point counts records in one month. Use it to spot busier and quieter periods; it does not imply causation."})
                src += len(period_counts) + 5

        # A histogram for one numeric field is more meaningful than comparing averages with different units.
        if useful_numeric:
            selected_num = max(useful_numeric, key=lambda c: work[c].notna().sum())
            vals = pd.to_numeric(work[selected_num], errors="coerce").dropna()
            if len(vals) >= 5 and vals.nunique() > 1:
                bins = min(10, max(5, int(np.sqrt(len(vals)))))
                try:
                    bucket = pd.cut(vals, bins=bins, duplicates="drop")
                    distribution = bucket.value_counts(sort=False)
                    distribution = distribution[distribution > 0]
                    if len(distribution) >= 2:
                        analysis.cell(src, 1, f"CHART DATA - DISTRIBUTION: {selected_num}")
                        analysis.cell(src + 1, 1, "Value range"); analysis.cell(src + 1, 2, "Record count")
                        for i, (interval, count) in enumerate(distribution.items(), src + 2):
                            analysis.cell(i, 1, str(interval)[:70]); analysis.cell(i, 2, int(count))
                        charts.append({"kind":"bar", "title":f"Distribution of {selected_num}", "header":src+1, "count":len(distribution), "color":blue,
                                       "note":f"Groups values in '{selected_num}' into ranges. Taller bars show which ranges occur more often; this is not a comparison with other columns."})
                        src += len(distribution) + 5
                except (ValueError, TypeError):
                    pass

        # Dashboard header and KPI cards
        dash.merge_cells("A1:P1"); dash["A1"] = "AI DATA INTELLIGENCE PRO  |  SMART DATA DASHBOARD"
        dash["A1"].font = Font(name="Arial", size=20, bold=True, color="FFFFFF")
        dash["A1"].fill = PatternFill("solid", fgColor=navy); dash["A1"].alignment = Alignment(vertical="center")
        dash.row_dimensions[1].height = 40
        dash.merge_cells("A2:P2"); dash["A2"] = "Dataset overview • Data quality • Trends • Distribution • Actionable findings"
        dash["A2"].font = Font(name="Arial", size=11, italic=True, color="53677D")
        dash.row_dimensions[2].height = 24
        metrics = [("TOTAL RECORDS", nrows), ("TOTAL COLUMNS", ncols), ("MISSING CELLS", missing_total),
                   ("DUPLICATE ROWS", duplicates), ("NUMERIC FIELDS", len(numeric_cols)), ("DATE FIELDS", len(date_cols))]
        for idx, (label, value) in enumerate(metrics):
            start_col = (idx % 3) * 5 + 1
            start_row = 4 + (idx // 3) * 3
            dash.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=start_col+3)
            dash.merge_cells(start_row=start_row+1, start_column=start_col, end_row=start_row+1, end_column=start_col+3)
            lc = dash.cell(start_row, start_col, label); lc.fill = PatternFill("solid", fgColor=teal)
            lc.font = Font(name="Arial", bold=True, size=10, color="FFFFFF"); lc.alignment = Alignment(horizontal="center", vertical="center")
            vc = dash.cell(start_row+1, start_col, value); vc.fill = PatternFill("solid", fgColor="EAF3F7")
            vc.font = Font(name="Arial", bold=True, size=19, color=navy); vc.alignment = Alignment(horizontal="center", vertical="center")
            for rr in (start_row, start_row+1):
                for cc in range(start_col, start_col+4):
                    dash.cell(rr, cc).border = Border(left=border, right=border, top=border, bottom=border)
            dash.row_dimensions[start_row].height = 22; dash.row_dimensions[start_row+1].height = 32

        findings = []
        findings.append(f"The dataset contains {nrows:,} records across {ncols:,} columns.")
        findings.append(f"Data completeness is {100-missing_pct:.1f}%; {missing_total:,} cells are blank." if missing_total else "Data completeness is 100%: no blank cells were detected.")
        findings.append(f"Found {duplicates:,} duplicate row(s)." if duplicates else "No fully duplicated rows were detected.")
        if candidates:
            top_col = max(candidates, key=lambda c: work[c].nunique(dropna=True))
            top_counts = work[top_col].fillna("(Missing)").astype(str).value_counts()
            if not top_counts.empty:
                findings.append(f"Most frequent value: '{top_counts.index[0]}' in '{top_col}' ({top_counts.iloc[0]:,} records, {top_counts.iloc[0]/max(nrows,1)*100:.1f}%).")
        if date_cols:
            findings.append(f"A date field ('{next(iter(date_cols))}') was detected, so a time trend is available.")
        if useful_numeric:
            findings.append(f"{len(useful_numeric)} numeric field(s) can be explored; see Analysis for averages, medians, ranges and distribution.")
        dash.merge_cells("A11:P11"); dash["A11"] = "WHAT THE DATA TELLS YOU"
        dash["A11"].fill = PatternFill("solid", fgColor=navy); dash["A11"].font = Font(bold=True, color="FFFFFF", size=12)
        dash.row_dimensions[11].height = 27
        for i, insight in enumerate(findings[:5], 12):
            dash.merge_cells(start_row=i, start_column=1, end_row=i, end_column=16)
            cell = dash.cell(i, 1, "• " + insight); cell.fill = PatternFill("solid", fgColor=pale if i % 2 == 0 else "FFFFFF")
            cell.font = Font(name="Arial", size=10, color="24364B"); cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=border); dash.row_dimensions[i].height = 30

        # Two-column chart grid with readable chart dimensions and restrained styling.
        anchors = [("A19", 19), ("I19", 19), ("A38", 38), ("I38", 38)]
        for idx, item in enumerate(charts[:4]):
            anchor, _ = anchors[idx]
            count = item["count"]
            if count < 1: continue
            chart = LineChart() if item["kind"] == "line" else BarChart()
            chart.style = 10
            chart.title = item["title"]
            chart.height = 8.2; chart.width = 15.5
            chart.legend = None
            chart.display_blanks = "gap"
            chart.y_axis.majorGridlines = None
            chart.x_axis.majorGridlines = None
            data_ref = Reference(analysis, min_col=2, min_row=item["header"], max_row=item["header"] + count)
            cats_ref = Reference(analysis, min_col=1, min_row=item["header"] + 1, max_row=item["header"] + count)
            chart.add_data(data_ref, titles_from_data=True); chart.set_categories(cats_ref)
            if item["kind"] == "bar":
                chart.type = "bar"; chart.grouping = "clustered"; chart.overlap = 0; chart.gapWidth = 55
                chart.x_axis.title = "Number of records" if "Distribution" not in item["title"] else "Records"
                chart.y_axis.title = "Group / range"
                chart.y_axis.scaling.orientation = "maxMin"
            else:
                chart.y_axis.title = "Number of records"; chart.x_axis.title = "Month"
                chart.height = 8.5; chart.width = 15.5
                chart.smooth = False; chart.marker = "circle"
            if chart.series:
                chart.series[0].graphicalProperties.solidFill = item["color"]
                chart.series[0].graphicalProperties.line.solidFill = item["color"]
                chart.dataLabels = DataLabelList(); chart.dataLabels.showVal = True
            dash.add_chart(chart, anchor)
            note_row = 35 if idx < 2 else 54
            note_col = 1 if idx % 2 == 0 else 9
            dash.merge_cells(start_row=note_row, start_column=note_col, end_row=note_row+1, end_column=note_col+6)
            note = dash.cell(note_row, note_col, "INSIGHT: " + item["note"])
            note.font = Font(name="Arial", size=9, italic=True, color="42566B")
            note.fill = PatternFill("solid", fgColor="EAF3F7")
            note.alignment = Alignment(wrap_text=True, vertical="center")
            dash.row_dimensions[note_row].height = 27; dash.row_dimensions[note_row+1].height = 22

        if not charts:
            dash.merge_cells("A19:P21")
            dash["A19"] = "Not enough suitable fields were found to create a meaningful chart. See the Analysis sheet for the available statistics."
            dash["A19"].alignment = Alignment(wrap_text=True, vertical="center")

        # Worksheet formatting
        for ws in (data_ws, analysis):
            ws.freeze_panes = "A2"
            if ws.max_row and ws.max_column:
                if ws.title == "Data": ws.auto_filter.ref = ws.dimensions
                for cell in ws[1]:
                    cell.font = Font(bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor=navy)
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                for col_cells in ws.columns:
                    letter = get_column_letter(col_cells[0].column)
                    max_len = max((len(str(c.value or "")) for c in list(col_cells)[:300]), default=0)
                    ws.column_dimensions[letter].width = min(max(max_len + 2, 12), 32)
                for row in ws.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                            cell.number_format = "#,##0.00"
                if ws.title == "Analysis":
                    for row in ws.iter_rows():
                        value = row[0].value
                        if isinstance(value, str) and value.isupper() and ("DATASET" in value or "STATISTICS" in value or "BREAKDOWNS" in value or "CHART DATA" in value):
                            row[0].fill = PatternFill("solid", fgColor=teal)
                            row[0].font = Font(bold=True, color="FFFFFF")
        for col in range(1, 17): dash.column_dimensions[get_column_letter(col)].width = 12
        dash.freeze_panes = "A4"
        dash.sheet_properties.pageSetUpPr.fitToPage = True
        dash.page_setup.fitToWidth = 1; dash.page_setup.fitToHeight = 0
        wb.active = 0
    return buf.getvalue()


def make_report(df, source_name):
    missing = df.isna().sum()
    duplicates = int(df.duplicated().sum())
    lines = [
        "AI DATA INTELLIGENCE PRO - DATA QUALITY REPORT",
        f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}",
        f"Source: {source_name}",
        f"Rows: {len(df)}",
        f"Columns: {len(df.columns)}",
        f"Duplicate rows: {duplicates}",
        f"Missing cells: {int(missing.sum())}",
        "",
        "COLUMN PROFILE"
    ]
    for c in df.columns:
        lines.append(f"- {c}: dtype={df[c].dtype}; missing={int(df[c].isna().sum())}; unique={int(df[c].nunique(dropna=True))}")
    lines += ["", "MISSING VALUES", missing.to_string(), "", "DESCRIPTIVE STATISTICS", df.describe(include="all").transpose().fillna("").to_string()]
    return "\n".join(lines)


def create_pdf_quality_report(df, source_name):
    """Build a readable, dataset-independent PDF quality report without changing source data."""
    buffer = io.BytesIO()
    df = df.copy(deep=True)
    rows, cols = df.shape
    total_cells = rows * cols
    missing = df.isna().sum()
    missing_total = int(missing.sum())
    duplicate_rows = int(df.duplicated().sum()) if rows else 0
    completeness = 100.0 * (total_cells - missing_total) / total_cells if total_cells else 100.0
    duplicate_free = 100.0 * (rows - duplicate_rows) / rows if rows else 100.0
    quality_score = 0.70 * completeness + 0.30 * duplicate_free

    # A separate sanitized copy is used only for statistics and charts.
    numeric = df.select_dtypes(include=[np.number]).copy()
    nonfinite_by_col = {}
    for column in numeric.columns:
        arr = pd.to_numeric(numeric[column], errors="coerce")
        count_nonfinite = int(np.isinf(arr.to_numpy(dtype=float, na_value=np.nan)).sum())
        if count_nonfinite:
            nonfinite_by_col[column] = count_nonfinite
        numeric[column] = arr.replace([np.inf, -np.inf], np.nan)

    palette = {
        "dark": colors.HexColor("#21140F"),
        "brown": colors.HexColor("#6F4E37"),
        "tan": colors.HexColor("#C49A6C"),
        "cream": colors.HexColor("#F4EEE7"),
        "pale": colors.HexColor("#FBF8F4"),
        "text": colors.HexColor("#30231D"),
        "muted": colors.HexColor("#75675D"),
        "line": colors.HexColor("#D8C8B8"),
        "green": colors.HexColor("#2E7D5B"),
        "red": colors.HexColor("#B54A3A"),
    }
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="AIReportTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=21, leading=25, textColor=colors.white, alignment=TA_CENTER, spaceAfter=4))
    styles.add(ParagraphStyle(name="AIReportSub", parent=styles["Normal"], fontName="Helvetica", fontSize=9, leading=12, textColor=colors.HexColor("#E6D8CB"), alignment=TA_CENTER))
    styles.add(ParagraphStyle(name="AISection", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=palette["brown"], spaceBefore=10, spaceAfter=5, keepWithNext=True))
    styles.add(ParagraphStyle(name="AIBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=8.5, leading=12, textColor=palette["text"], spaceAfter=3))
    styles.add(ParagraphStyle(name="AISmall", parent=styles["BodyText"], fontName="Helvetica", fontSize=7, leading=9, textColor=palette["text"], wordWrap="CJK"))
    styles.add(ParagraphStyle(name="AITableHead", parent=styles["AISmall"], fontName="Helvetica-Bold", textColor=colors.white, leading=9))
    styles.add(ParagraphStyle(name="AIKpiLabel", parent=styles["AISmall"], fontName="Helvetica-Bold", fontSize=6.8, alignment=TA_CENTER, textColor=colors.white))
    styles.add(ParagraphStyle(name="AIKpiValue", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=13, leading=16, alignment=TA_CENTER, textColor=palette["dark"]))

    def txt(value, limit=120):
        try:
            value = str(value).replace("\n", " ").replace("\r", " ")
        except Exception:
            value = "N/A"
        if len(value) > limit:
            value = value[:limit - 3] + "..."
        return escape(value)

    def fmt(value, digits=2):
        try:
            if pd.isna(value) or not np.isfinite(float(value)):
                return "N/A"
            return f"{float(value):,.{digits}f}"
        except Exception:
            return "N/A"

    def styled_table(data, widths, font_size=7):
        table = LongTable(data, colWidths=widths, repeatRows=1, hAlign="LEFT", splitByRow=1)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), palette["dark"]),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("GRID", (0, 0), (-1, -1), 0.35, palette["line"]),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, palette["pale"]]),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ]))
        return table

    page_w, page_h = landscape(A4)
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), rightMargin=12*mm, leftMargin=12*mm,
                            topMargin=12*mm, bottomMargin=15*mm,
                            title="AI Data Intelligence Pro - Data Quality Report", author="AI Data Intelligence Pro")
    content_w = page_w - 24*mm
    story = []
    header = Table([
        [Paragraph("AI DATA INTELLIGENCE PRO", styles["AIReportTitle"])],
        [Paragraph("DATA QUALITY REPORT  |  CLEAR SUMMARY & ACTIONABLE FINDINGS", styles["AIReportSub"])],
    ], colWidths=[content_w])
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), palette["dark"]),
        ("BOX", (0, 0), (-1, -1), 1, palette["tan"]),
        ("TOPPADDING", (0, 0), (-1, 0), 15), ("BOTTOMPADDING", (0, 0), (-1, 0), 3),
        ("TOPPADDING", (0, 1), (-1, 1), 2), ("BOTTOMPADDING", (0, 1), (-1, 1), 14),
    ]))
    story += [header, Spacer(1, 5*mm)]
    story.append(Paragraph(f"<b>Dataset:</b> {txt(source_name, 180)}", styles["AIBody"]))
    story.append(Paragraph(f"<b>Generated:</b> {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", styles["AIBody"]))
    story.append(Spacer(1, 3*mm))

    kpi_items = [
        ("TOTAL RECORDS", f"{rows:,}"), ("TOTAL COLUMNS", f"{cols:,}"),
        ("MISSING CELLS", f"{missing_total:,}"), ("DUPLICATE ROWS", f"{duplicate_rows:,}"),
        ("COMPLETENESS", f"{completeness:.1f}%")
    ]
    kpi_data = [[Paragraph(label, styles["AIKpiLabel"]) for label, _ in kpi_items],
                [Paragraph(value, styles["AIKpiValue"]) for _, value in kpi_items]]
    kpi = Table(kpi_data, colWidths=[content_w/5]*5, rowHeights=[9*mm, 13*mm])
    kpi.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), palette["brown"]), ("BACKGROUND", (0, 1), (-1, 1), palette["cream"]),
        ("BOX", (0, 0), (-1, -1), 0.8, palette["tan"]), ("INNERGRID", (0, 0), (-1, -1), 0.4, palette["line"]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ALIGN", (0, 0), (-1, -1), "CENTER"),
        ("LEFTPADDING", (0, 0), (-1, -1), 3), ("RIGHTPADDING", (0, 0), (-1, -1), 3),
    ]))
    story += [kpi, Spacer(1, 4*mm)]

    score_color = palette["green"] if quality_score >= 95 else (palette["brown"] if quality_score >= 80 else palette["red"])
    score_box = Table([[Paragraph(f"INDICATIVE QUALITY SCORE: <b>{quality_score:.2f} / 100</b>", styles["AIBody"])],
                       [Paragraph("Calculated as 70% completeness + 30% duplicate-free row rate. This is a limited indicator only; it does not verify factual accuracy, business rules, or real-world correctness.", styles["AISmall"])]], colWidths=[content_w])
    score_box.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), palette["pale"]), ("BOX", (0, 0), (-1, -1), 0.8, score_color),
        ("LINEBELOW", (0, 0), (-1, 0), 0.4, palette["line"]), ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story += [score_box, Spacer(1, 2*mm)]

    story.append(Paragraph("1. Quick Summary", styles["AISection"]))
    if rows == 0:
        summary = "The dataset has column headers but no records. Import records before interpreting statistics or building models."
    else:
        summary = f"The dataset contains {rows:,} records across {cols:,} columns. "
        summary += (f"There are {missing_total:,} missing cells ({100*missing_total/total_cells:.2f}% of all cells). " if total_cells else "")
        summary += f"There are {duplicate_rows:,} exact duplicate rows. "
        summary += "No missing cells were detected." if missing_total == 0 else "Missing values should be reviewed column by column before modeling."
    story.append(Paragraph(txt(summary, 800), styles["AIBody"]))
    if nonfinite_by_col:
        total_nonfinite = sum(nonfinite_by_col.values())
        story.append(Paragraph(f"<b>Important numeric issue:</b> {total_nonfinite:,} infinite (+/-inf) values were found across {len(nonfinite_by_col)} numeric column(s). They are excluded from the report's numeric summaries and charts; the original dataset is not modified.", styles["AIBody"]))

    story.append(Paragraph("2. Missing Values", styles["AISection"]))
    missing_cols = missing[missing > 0].sort_values(ascending=False).head(15)
    if len(missing_cols):
        chart_buf = io.BytesIO()
        fig, ax = plt.subplots(figsize=(9.6, max(2.2, min(4.0, 0.28*len(missing_cols)+0.7))))
        ax.barh([str(c)[:34] for c in missing_cols.index][::-1], missing_cols.values[::-1], color="#8B5E3C")
        ax.set_title("Columns with the Most Missing Values", fontsize=11, pad=10)
        ax.set_xlabel("Missing cells")
        ax.grid(axis="x", alpha=0.2)
        fig.tight_layout()
        fig.savefig(chart_buf, format="png", dpi=140, bbox_inches="tight")
        plt.close(fig)
        chart_buf.seek(0)
        story.append(Image(chart_buf, width=min(content_w, 235*mm), height=min(78*mm, max(48*mm, 0.28*len(missing_cols)*mm*10))))
        miss_data = [[Paragraph("Column", styles["AITableHead"]), Paragraph("Missing Cells", styles["AITableHead"]), Paragraph("Missing %", styles["AITableHead"])]]
        for c, count in missing_cols.items():
            miss_data.append([Paragraph(txt(c, 100), styles["AISmall"]), f"{int(count):,}", f"{(count/rows*100 if rows else 0):.2f}%"])
        story.append(styled_table(miss_data, [content_w*0.58, content_w*0.21, content_w*0.21]))
    else:
        story.append(Paragraph("<b>No missing values detected.</b> Every cell in the loaded dataset is populated. This does not guarantee that values are correct or meaningful.", styles["AIBody"]))

    story.append(Paragraph("3. Column Profile", styles["AISection"]))
    profile = [[Paragraph(h, styles["AITableHead"]) for h in ["Column", "Data Type", "Non-null", "Missing", "Missing %", "Unique Values", "Infinite Values"]]]
    for c in df.columns:
        try:
            unique = f"{int(df[c].nunique(dropna=True)):,}"
        except Exception:
            unique = "N/A"
        miss = int(df[c].isna().sum())
        nonnull = int(df[c].notna().sum())
        inf_count = nonfinite_by_col.get(c, 0)
        profile.append([
            Paragraph(txt(c, 100), styles["AISmall"]), Paragraph(txt(df[c].dtype, 30), styles["AISmall"]),
            f"{nonnull:,}", f"{miss:,}", f"{(miss/rows*100 if rows else 0):.2f}%", unique, f"{inf_count:,}" if inf_count else "0"
        ])
    story.append(styled_table(profile, [content_w*.24, content_w*.15, content_w*.12, content_w*.10, content_w*.12, content_w*.14, content_w*.13]))
    unclear_names = [str(c) for c in df.columns if str(c).strip() in {"#NAME?", "Unnamed: 0", "Unnamed"}]
    if unclear_names:
        story.append(Spacer(1, 2*mm))
        story.append(Paragraph(f"<b>Column name to review:</b> {txt(', '.join(unclear_names), 200)}. The report does not rename or delete it; confirm whether it represents a valid field or an import/formula artifact.", styles["AIBody"]))

    if not numeric.empty:
        story.append(Paragraph("4. Numeric Statistics", styles["AISection"]))
        stat_data = [[Paragraph(h, styles["AITableHead"]) for h in ["Column", "Valid Count", "Mean", "Std Dev", "Minimum", "Median", "Maximum"]]]
        for c in numeric.columns:
            s = numeric[c].dropna()
            stat_data.append([
                Paragraph(txt(c, 90), styles["AISmall"]), f"{len(s):,}", fmt(s.mean()), fmt(s.std()),
                fmt(s.min()), fmt(s.median()), fmt(s.max())
            ])
        story.append(styled_table(stat_data, [content_w*.20, content_w*.12, content_w*.13, content_w*.13, content_w*.13, content_w*.13, content_w*.16]))
        # Choose a numeric column with at least two finite values for a meaningful histogram.
        chart_col = next((c for c in numeric.columns if numeric[c].notna().sum() > 1 and numeric[c].nunique(dropna=True) > 1), None)
        if chart_col is not None:
            vals = numeric[chart_col].dropna()
            chart_buf = io.BytesIO()
            fig, ax = plt.subplots(figsize=(9.5, 3.0))
            ax.hist(vals, bins=min(20, max(5, int(np.sqrt(len(vals))))), color="#C49A6C", edgecolor="white")
            ax.set_title(f"Distribution of {str(chart_col)[:60]}", fontsize=11, pad=8)
            ax.set_xlabel(str(chart_col)[:60]); ax.set_ylabel("Frequency")
            ax.grid(axis="y", alpha=0.2)
            fig.tight_layout()
            fig.savefig(chart_buf, format="png", dpi=140, bbox_inches="tight")
            plt.close(fig); chart_buf.seek(0)
            story += [Spacer(1, 3*mm), Image(chart_buf, width=min(content_w, 220*mm), height=67*mm)]
    else:
        story.append(Paragraph("No numeric columns were detected, so numeric statistics and distribution charts are not applicable to this dataset.", styles["AIBody"]))

    story.append(Paragraph("5. Recommended Actions", styles["AISection"]))
    recommendations = []
    if missing_total:
        worst_col = str(missing_cols.index[0]) if len(missing_cols) else "the affected columns"
        recommendations.append(f"Review missing values first, starting with {worst_col}. Choose fill, removal, or retention based on the field's meaning and downstream use.")
    else:
        recommendations.append("Completeness is currently 100%. Keep checking new imports because future files may introduce missing values.")
    if duplicate_rows:
        recommendations.append(f"Inspect {duplicate_rows:,} exact duplicate rows before removing them; repeated rows may be legitimate in some datasets.")
    else:
        recommendations.append("No exact duplicate rows were detected by comparing complete rows.")
    if nonfinite_by_col:
        recommendations.append("Investigate infinite numeric values and determine why they occur (for example, division by zero). They were excluded from these statistics, not altered in the source data.")
    if unclear_names:
        recommendations.append("Confirm the meaning of the unclear column name(s) listed above before relying on them in analysis.")
    if rows == 0:
        recommendations.append("The dataset contains no records. Load a non-empty file before drawing conclusions.")
    recommendations.append("Validate domain-specific rules and real-world accuracy separately; a high score here is not proof that the data is correct.")
    for i, item in enumerate(recommendations, 1):
        story.append(Paragraph(f"<b>{i}.</b> {txt(item, 700)}", styles["AIBody"]))
        story.append(Spacer(1, 1.5*mm))

    story.append(Spacer(1, 4*mm))
    story.append(Paragraph("Generated by AI Data Intelligence Pro • Report metrics describe the loaded dataset only.", styles["AISmall"]))

    def page_number(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(palette["line"])
        canvas.line(12*mm, 11*mm, page_w-12*mm, 11*mm)
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(palette["muted"])
        canvas.drawString(12*mm, 7*mm, "AI Data Intelligence Pro | Data Quality Report")
        canvas.drawRightString(page_w-12*mm, 7*mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=page_number, onLaterPages=page_number)
    buffer.seek(0)
    return buffer.getvalue()

with st.sidebar:
    st.markdown("## 🧠 AI Intelligence")
    st.markdown('<span class="badge">PRO • DATA LAB</span>', unsafe_allow_html=True)
    st.markdown("---")
    uploaded = st.file_uploader("حمّل Dataset", type=["csv", "xlsx", "xls", "json", "parquet"], help="الحد المضبوط في إعدادات التطبيق يصل إلى 500 MB. ملفات Excel لها حدود داخلية حسب صيغة الملف والذاكرة المتاحة.")
    if uploaded is not None:
        upload_signature = f"{getattr(uploaded, 'file_id', '')}:{uploaded.name}:{uploaded.size}"
        if upload_signature != st.session_state.get("loaded_upload_signature"):
            try:
                with st.spinner("جاري قراءة الملف بطريقة موفرة للذاكرة... الملفات الكبيرة قد تستغرق وقتًا أطول"):
                    new_df = read_dataset(uploaded)
                if new_df.empty or len(new_df.columns) == 0:
                    st.error("الملف لا يحتوي على بيانات قابلة للقراءة.")
                else:
                    st.session_state.df = new_df
                    st.session_state.source_name = uploaded.name
                    st.session_state.clean_df = None
                    st.session_state.trained_model = None
                    st.session_state.model_info = None
                    st.session_state.automl_results = None
                    st.session_state.outlier_df = None
                    st.session_state.cluster_df = None
                    for export_key in ("export_pdf_bytes", "export_excel_bytes", "export_csv_bytes", "export_json_bytes"):
                        st.session_state.pop(export_key, None)
                    gc.collect()
                    st.session_state.loaded_upload_signature = upload_signature
                    st.success(f"تم تحميل {uploaded.name}")
            except Exception as e:
                st.error(f"تعذر قراءة الملف: {e}")
    if st.session_state.df is not None:
        st.markdown("---")
        st.caption(f"الملف الحالي: {st.session_state.source_name}")
        st.caption(f"{len(st.session_state.df):,} صف × {len(st.session_state.df.columns):,} عمود")
        if st.button("إزالة الملف الحالي", width="stretch"):
            for k in ["df", "clean_df", "trained_model", "model_info", "automl_results", "outlier_df", "cluster_df"]:
                st.session_state[k] = None
            st.session_state.source_name = ""
            st.rerun()
    st.markdown("---")
    st.markdown('<div class="small-muted">تحليل محلي • لا يتم إرسال ملفاتك إلى خدمة خارجية</div>', unsafe_allow_html=True)

st.markdown("""
<div class="hero">
  <div class="hero-title">AI Data Intelligence & Automation</div>
  <p class="hero-sub">منصة تحليل بيانات ذكية: استكشاف • تنظيف • AutoML • اكتشاف الشذوذ • تقارير • SQL</p>
</div>
""", unsafe_allow_html=True)

df = st.session_state.df
if df is None:
    st.markdown("""
    <div class="section-card">
      <h3>ابدأ بتحليل بياناتك</h3>
      <p>حمّل ملف CSV أو Excel أو JSON أو Parquet من القائمة الجانبية. سيعرض النظام خصائص البيانات ويتيح لك تنظيفها وتجربة نماذج تعلم الآلة.</p>
      <p><b>ملاحظة:</b> نتائج النماذج تعتمد على جودة البيانات واختيار عمود الهدف، ولا يوجد نموذج واحد مناسب لكل مجموعة بيانات.</p>
    </div>
    """, unsafe_allow_html=True)
    demo = pd.DataFrame({
        "Age": [22,35,29,41,27,33,48,31,26,39,44,24],
        "Department": ["IT","Sales","IT","HR","Sales","HR","IT","Sales","HR","IT","Sales","HR"],
        "Monthly_Spend": [1200,1800,1500,900,2100,1100,2300,1700,1000,2000,2200,950],
        "Customer_Type": ["New","Returning","New","New","Returning","New","Returning","Returning","New","Returning","Returning","New"],
        "Churn": ["No","No","Yes","No","Yes","No","No","Yes","No","No","Yes","No"]
    })
    st.markdown("#### معاينة توضيحية")
    st.dataframe(demo, width="stretch", hide_index=True)
    st.info("ارفع ملفك من القائمة الجانبية لفتح أدوات التحليل الفعلية.")
    st.stop()

df_original = df
active_df = st.session_state.clean_df if st.session_state.clean_df is not None else df_original
large_dataset = len(active_df) > 200_000
chart_df = plot_sample(active_df, 50_000)

tabs = st.tabs(["🏠 نظرة عامة", "🔎 التحليل التلقائي", "🧹 جودة وتنظيف", "🤖 Machine Learning", "🧬 Clustering", "🚨 Anomaly Detection", "📑 التقارير والتصدير", "🗄️ SQL & Automation"])

with tabs[0]:
    st.markdown("### ملخص مجموعة البيانات")
    c1,c2,c3,c4,c5 = st.columns(5)
    missing_total = int(active_df.isna().sum().sum())
    duplicate_count = safe_duplicate_count(active_df)
    with c1: metric_card("عدد الصفوف", f"{len(active_df):,}", "Records")
    with c2: metric_card("عدد الأعمدة", f"{len(active_df.columns):,}", "Features")
    with c3: metric_card("قيم مفقودة", f"{missing_total:,}", "Missing cells")
    with c4: metric_card("صفوف مكررة", f"{duplicate_count:,}" if duplicate_count is not None else "تخطي", "Exact check skipped above 200k rows")
    with c5: metric_card("استهلاك الذاكرة", f"{active_df.memory_usage(deep=True).sum()/1024**2:.2f} MB", "Approx.")
    st.markdown("### معاينة البيانات")
    st.dataframe(active_df.head(100), width="stretch", hide_index=True)
    left, right = st.columns(2)
    with left:
        st.markdown("### أنواع الأعمدة")
        types_df = active_df.dtypes.astype(str).value_counts().rename_axis("Data Type").reset_index(name="Columns")
        fig = px.bar(types_df, x="Data Type", y="Columns", color="Data Type", template="plotly_dark")
        fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", showlegend=False)
        st.plotly_chart(fig, width="stretch")
    with right:
        st.markdown("### القيم المفقودة حسب العمود")
        miss_df = active_df.isna().sum().sort_values(ascending=False).head(20).reset_index()
        miss_df.columns = ["Column", "Missing"]
        fig = px.bar(miss_df, x="Missing", y="Column", orientation="h", template="plotly_dark")
        fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", yaxis={"categoryorder":"total ascending"})
        st.plotly_chart(fig, width="stretch")
    st.markdown("### الملف الشخصي للأعمدة")
    profile = pd.DataFrame({
        "Column": active_df.columns,
        "Type": [str(active_df[c].dtype) for c in active_df.columns],
        "Non-null": [int(active_df[c].notna().sum()) for c in active_df.columns],
        "Missing": [int(active_df[c].isna().sum()) for c in active_df.columns],
        "Missing %": [round(active_df[c].isna().mean()*100, 2) for c in active_df.columns],
        "Unique": []
    })
    unique_values = [safe_unique_count(active_df[c]) for c in active_df.columns]
    profile["Unique"] = [("~" if sampled else "") + f"{count:,}" for count, sampled in unique_values]
    st.dataframe(profile, width="stretch", hide_index=True)

with tabs[1]:
    st.markdown("### تحليل تلقائي استكشافي")
    st.caption("يعرض النظام اقتراحات ورسومًا أولية اعتمادًا على أنواع الأعمدة. راجع النتائج قبل اتخاذ قرار.")
    num_cols = active_df.select_dtypes(include=np.number).columns.tolist()
    cat_cols = [c for c in active_df.columns if c not in num_cols]
    if num_cols:
        col = st.selectbox("اختر عمودًا رقميًا", num_cols, key="auto_num_col")
        vals = pd.to_numeric(active_df[col], errors="coerce").dropna()
        a,b,c,d = st.columns(4)
        with a: metric_card("المتوسط", f"{vals.mean():,.3f}" if len(vals) else "—")
        with b: metric_card("الوسيط", f"{vals.median():,.3f}" if len(vals) else "—")
        with c: metric_card("الانحراف المعياري", f"{vals.std():,.3f}" if len(vals) else "—")
        with d: metric_card("القيم الفريدة", f"{vals.nunique():,}")
        f1,f2 = st.columns(2)
        with f1:
            fig = px.histogram(chart_df, x=col, nbins=30, title=f"Distribution: {col} (sampled for large files)", template="plotly_dark")
            fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
            st.plotly_chart(fig, width="stretch")
        with f2:
            fig = px.box(chart_df, y=col, title=f"Outlier overview: {col} (sampled for large files)", template="plotly_dark")
            fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
            st.plotly_chart(fig, width="stretch")
    if cat_cols:
        col_cat = st.selectbox("اختر عمودًا فئويًا", cat_cols, key="auto_cat_col")
        counts = chart_df[col_cat].astype("string").fillna("(Missing)").value_counts().head(20).reset_index()
        counts.columns = [col_cat, "Count"]
        fig = px.bar(counts, x=col_cat, y="Count", title=f"Top categories: {col_cat}", template="plotly_dark")
        fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)", xaxis_tickangle=-25)
        st.plotly_chart(fig, width="stretch")
    if len(num_cols) >= 2:
        st.markdown("### ارتباط الأعمدة الرقمية")
        corr = chart_df[num_cols].corr(numeric_only=True)
        fig = px.imshow(corr, text_auto=".2f", aspect="auto", color_continuous_scale="Tealrose", template="plotly_dark")
        fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
        st.plotly_chart(fig, width="stretch")

with tabs[2]:
    st.markdown("### جودة البيانات والتنظيف")
    st.warning("التنظيف لا يغيّر الملف الأصلي؛ عند التطبيق تُحفظ نسخة منفصلة داخل الجلسة.")
    q1,q2,q3 = st.columns(3)
    with q1: metric_card("Missing cells", f"{int(active_df.isna().sum().sum()):,}")
    quality_duplicates = safe_duplicate_count(active_df)
    with q2: metric_card("Duplicate rows", f"{quality_duplicates:,}" if quality_duplicates is not None else "تخطي", "Exact check skipped above 200k rows")
    with q3: metric_card("Columns with missing", f"{int((active_df.isna().sum()>0).sum()):,}")
    with st.form("clean_form"):
        drop_duplicates = st.checkbox("حذف الصفوف المكررة", value=True)
        fill_numeric = st.selectbox("القيم المفقودة الرقمية", ["Median", "Mean", "Leave unchanged"], index=0)
        fill_categorical = st.selectbox("القيم المفقودة النصية", ["Most frequent", "Fill with Unknown", "Leave unchanged"], index=0)
        trim_text = st.checkbox("إزالة المسافات الزائدة من النصوص", value=True)
        normalize_headers = st.checkbox("تنسيق أسماء الأعمدة (اختياري)", value=False)
        submitted = st.form_submit_button("🧹 إنشاء نسخة نظيفة", width="stretch")
    if submitted:
        cleaned = active_df.copy()
        if trim_text:
            for c in cleaned.select_dtypes(include=["object", "string"]).columns:
                cleaned[c] = cleaned[c].map(lambda x: x.strip() if isinstance(x, str) else x)
        if drop_duplicates and len(cleaned) <= 200_000:
            cleaned = cleaned.drop_duplicates()
        elif drop_duplicates:
            st.warning("تم تخطي حذف التكرارات تلقائيًا لأن الملف أكبر من 200,000 صف لتجنب انهيار الذاكرة. يمكنك تنظيف نسخة أصغر أو استخدام أداة معالجة خارجية للملفات الضخمة.")
        for c in cleaned.columns:
            if pd.api.types.is_numeric_dtype(cleaned[c]):
                if fill_numeric == "Median" and cleaned[c].isna().any():
                    cleaned[c] = cleaned[c].fillna(cleaned[c].median())
                elif fill_numeric == "Mean" and cleaned[c].isna().any():
                    cleaned[c] = cleaned[c].fillna(cleaned[c].mean())
            elif fill_categorical == "Most frequent" and cleaned[c].isna().any() and not cleaned[c].dropna().empty:
                cleaned[c] = cleaned[c].fillna(cleaned[c].mode(dropna=True).iloc[0])
            elif fill_categorical == "Fill with Unknown":
                cleaned[c] = cleaned[c].fillna("Unknown")
        if normalize_headers:
            cleaned.columns = [str(c).strip().replace(" ", "_") for c in cleaned.columns]
        st.session_state.clean_df = cleaned
        st.success(f"تم إنشاء النسخة النظيفة: {len(cleaned):,} صف × {len(cleaned.columns):,} عمود")
    if st.session_state.clean_df is not None:
        st.markdown("### معاينة النسخة النظيفة")
        st.dataframe(st.session_state.clean_df.head(100), width="stretch", hide_index=True)
        d1,d2 = st.columns(2)
        with d1:
            st.download_button("تنزيل CSV النظيف", to_csv_bytes(st.session_state.clean_df), file_name="cleaned_dataset.csv", mime="text/csv", width="stretch")
        with d2:
            st.download_button("تنزيل Excel النظيف", to_excel_bytes(st.session_state.clean_df), file_name="cleaned_dataset.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
        if st.button("إلغاء النسخة النظيفة والعودة للأصل"):
            st.session_state.clean_df = None
            st.rerun()

with tabs[3]:
    st.markdown("### AutoML — تدريب ومقارنة النماذج")
    st.write("اختر عمود الهدف. سيقوم النظام بتجهيز الأعمدة الرقمية والنصية، وتعويض القيم المفقودة، وترميز الفئات داخل Pipeline لتقليل أخطاء التدريب.")
    ml_df = active_df
    candidates = infer_target_candidates(ml_df)
    if len(ml_df) < 10:
        st.error("عدد الصفوف قليل جدًا. يُفضّل توفير 10 صفوف على الأقل، ويفضل أكثر للحصول على تقييم موثوق.")
    elif not candidates:
        st.error("لا يوجد عمود هدف مناسب. يلزم عمود يحتوي على قيمتين مختلفتين على الأقل وبيانات كافية.")
    else:
        target = st.selectbox("عمود الهدف (Target)", candidates, key="ml_target")
        suggested = classify_task(ml_df[target])
        task = st.radio("نوع المهمة", ["Auto detect", "Classification", "Regression"], horizontal=True, index=0, key="ml_task")
        actual_task = suggested if task == "Auto detect" else task
        st.info(f"المهمة المقترحة: **{suggested}** • المهمة المستخدمة: **{actual_task}**")
        feature_options = [c for c in ml_df.columns if c != target]
        selected_features = st.multiselect("الأعمدة المستخدمة كمدخلات (Features)", feature_options, default=feature_options[:min(len(feature_options), 30)], key="ml_features")
        test_size = st.slider("نسبة بيانات الاختبار", min_value=0.1, max_value=0.4, value=0.2, step=0.05, key="ml_test_size")
        run_automl = st.button("🚀 تدريب ومقارنة النماذج", width="stretch")
        if run_automl:
            if not selected_features:
                st.error("اختر عمودًا واحدًا على الأقل كمدخل.")
            else:
                work = ml_df[selected_features + [target]].copy()
                work = work.replace([np.inf, -np.inf], np.nan)
                work = work.loc[work[target].notna()].copy()
                y = work[target].copy()
                X = work[selected_features].copy()
                # Remove columns that have no useful values
                X = X.loc[:, X.notna().any(axis=0)]
                X = X.loc[:, X.nunique(dropna=True) > 1]
                if X.shape[1] == 0:
                    st.error("الأعمدة المختارة لا تحتوي على اختلافات مفيدة أو كلها فارغة.")
                elif len(work) < 10:
                    st.error(f"عدد الصفوف الصالحة بعد استبعاد القيم المفقودة في الهدف هو {len(work)} فقط. تحتاج إلى 10 صفوف على الأقل.")
                elif y.nunique(dropna=True) < 2:
                    st.error("عمود الهدف يحتوي على فئة/قيمة واحدة فقط. اختر عمود هدف مختلفًا.")
                else:
                    try:
                        if actual_task == "Classification":
                            y = y.astype(str)
                            class_counts = y.value_counts()
                            stratify = y if class_counts.min() >= 2 and len(class_counts) > 1 else None
                            Xtr, Xte, ytr, yte = train_test_split(X, y, test_size=test_size, random_state=42, stratify=stratify)
                            if ytr.nunique() < 2:
                                raise ValueError("مجموعة التدريب تحتوي على فئة واحدة فقط. اختر هدفًا آخر أو زِد حجم البيانات.")
                            pre = make_preprocessor(Xtr)
                            model_specs = [
                                ("Random Forest", RandomForestClassifier(n_estimators=100, random_state=42, class_weight="balanced", n_jobs=1)),
                                ("Extra Trees", ExtraTreesClassifier(n_estimators=100, random_state=42, class_weight="balanced", n_jobs=1))
                            ]
                            results, fitted = [], []
                            for name, estimator in model_specs:
                                pipe = Pipeline([("prep", pre), ("model", estimator)])
                                pipe.fit(Xtr, ytr)
                                pred = pipe.predict(Xte)
                                results.append({"Model": name, "Accuracy": accuracy_score(yte, pred), "Precision (weighted)": precision_score(yte, pred, average="weighted", zero_division=0), "Recall (weighted)": recall_score(yte, pred, average="weighted", zero_division=0), "F1 (weighted)": f1_score(yte, pred, average="weighted", zero_division=0)})
                                fitted.append((name, pipe, pred))
                            result_df = pd.DataFrame(results).sort_values("F1 (weighted)", ascending=False)
                            best_name = result_df.iloc[0]["Model"]
                            best_pipe = next(p for n,p,_ in fitted if n == best_name)
                            best_pred = next(pr for n,_,pr in fitted if n == best_name)
                            st.session_state.trained_model = best_pipe
                            st.session_state.model_info = {"task": actual_task, "target": target, "features": list(X.columns), "model": best_name, "score": float(result_df.iloc[0]["F1 (weighted)"]), "classes": sorted(y.unique().tolist())}
                            st.session_state.automl_results = result_df
                            st.success(f"اكتمل التدريب. أفضل نموذج: {best_name}")
                            st.dataframe(result_df.style.format({c:"{:.3f}" for c in result_df.columns if c != "Model"}), width="stretch", hide_index=True)
                            labels = sorted(pd.Series(yte).astype(str).unique().tolist())
                            cm = confusion_matrix(yte, best_pred, labels=labels)
                            fig = px.imshow(cm, x=labels, y=labels, text_auto=True, title=f"Confusion Matrix — {best_name}", labels=dict(x="Predicted", y="Actual", color="Count"), template="plotly_dark", color_continuous_scale="Teal")
                            fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
                            st.plotly_chart(fig, width="stretch")
                            st.markdown("#### تقرير التصنيف")
                            st.code(classification_report(yte, best_pred, zero_division=0))
                        else:
                            y = pd.to_numeric(y, errors="coerce")
                            valid = y.notna() & np.isfinite(y)
                            X, y = X.loc[valid].copy(), y.loc[valid].copy()
                            if len(y) < 10 or y.nunique() < 2:
                                st.error(f"عدد الصفوف الرقمية الصالحة: {len(y)}. اختر هدفًا رقميًا يحتوي على 10 صفوف صالحة على الأقل وقيمتين مختلفتين.")
                            else:
                                Xtr, Xte, ytr, yte = train_test_split(X, y, test_size=test_size, random_state=42)
                                pre = make_preprocessor(Xtr)
                                model_specs = [
                                    ("Random Forest", RandomForestRegressor(n_estimators=100, random_state=42, n_jobs=1)),
                                    ("Extra Trees", ExtraTreesRegressor(n_estimators=100, random_state=42, n_jobs=1))
                                ]
                                results, fitted = [], []
                                for name, estimator in model_specs:
                                    pipe = Pipeline([("prep", pre), ("model", estimator)])
                                    pipe.fit(Xtr, ytr)
                                    pred = pipe.predict(Xte)
                                    results.append({"Model": name, "MAE": mean_absolute_error(yte, pred), "RMSE": float(np.sqrt(mean_squared_error(yte, pred))), "R2": r2_score(yte, pred) if len(yte) > 1 else np.nan})
                                    fitted.append((name, pipe, pred))
                                result_df = pd.DataFrame(results).sort_values("RMSE", ascending=True)
                                best_name = result_df.iloc[0]["Model"]
                                best_pipe = next(p for n,p,_ in fitted if n == best_name)
                                best_pred = next(pr for n,_,pr in fitted if n == best_name)
                                st.session_state.trained_model = best_pipe
                                st.session_state.model_info = {"task": actual_task, "target": target, "features": list(X.columns), "model": best_name, "score": float(result_df.iloc[0]["RMSE"])}
                                st.session_state.automl_results = result_df
                                st.success(f"اكتمل التدريب. أفضل نموذج حسب RMSE: {best_name}")
                                st.dataframe(result_df.style.format({c:"{:.3f}" for c in result_df.columns if c != "Model"}), width="stretch", hide_index=True)
                                compare = pd.DataFrame({"Actual": yte.to_numpy(), "Predicted": best_pred})
                                fig = px.scatter(compare, x="Actual", y="Predicted", title=f"Actual vs Predicted — {best_name}", template="plotly_dark", trendline="ols")
                                fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
                                st.plotly_chart(fig, width="stretch")
                    except Exception as e:
                        st.error(f"تعذر تدريب النموذج: {type(e).__name__}: {e}")
                        st.info("جرّب هدفًا آخر، أو قلل الأعمدة المختارة، أو تأكد أن البيانات تحتوي على صفوف كافية.")
    if st.session_state.model_info:
        st.markdown("### آخر نموذج تم تدريبه")
        st.json(st.session_state.model_info)
        st.caption("النموذج محفوظ مؤقتًا خلال جلسة Streamlit الحالية فقط.")

with tabs[4]:
    st.markdown("### Clustering — اكتشاف مجموعات دون Target")
    st.write("يستخدم K-Means لتجميع السجلات المتشابهة. اختر أعمدة رقمية مناسبة، وراجع المجموعات قبل تفسيرها.")
    numeric_cols = active_df.select_dtypes(include=np.number).columns.tolist()
    if len(numeric_cols) < 2 or len(active_df) < 10:
        st.warning("تحتاج هذه الأداة إلى عمودين رقميين على الأقل و10 صفوف.")
    else:
        cluster_features = st.multiselect("الأعمدة الرقمية", numeric_cols, default=numeric_cols[:min(4,len(numeric_cols))], key="cluster_features")
        k = st.slider("عدد المجموعات (K)", 2, min(10, max(2, len(active_df)-1)), 3, key="cluster_k")
        if st.button("اكتشاف المجموعات", width="stretch"):
            if len(cluster_features) < 2:
                st.error("اختر عمودين رقميين على الأقل.")
            else:
                cluster_source = active_df if len(active_df) <= 100_000 else active_df.sample(n=100_000, random_state=42)
                temp = cluster_source[cluster_features].replace([np.inf,-np.inf], np.nan).copy()
                temp = temp.fillna(temp.median(numeric_only=True)).dropna()
                if len(temp) <= k:
                    st.error("عدد الصفوف الصالحة يجب أن يكون أكبر من K.")
                else:
                    scaler = StandardScaler()
                    scaled = scaler.fit_transform(temp)
                    labels = KMeans(n_clusters=k, random_state=42, n_init=10).fit_predict(scaled)
                    result = temp.copy()
                    result["Cluster"] = labels.astype(str)
                    st.session_state.cluster_df = result
                    score = silhouette_score(scaled, labels) if len(set(labels)) > 1 and len(temp) > k else np.nan
                    a,b = st.columns(2)
                    with a: metric_card("عدد المجموعات", str(k), "K-Means")
                    with b: metric_card("Silhouette score", f"{score:.3f}" if np.isfinite(score) else "—", "Higher is usually better")
                    if len(cluster_features) >= 2:
                        fig = px.scatter(result, x=cluster_features[0], y=cluster_features[1], color="Cluster", title="Cluster map", template="plotly_dark")
                        fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
                        st.plotly_chart(fig, width="stretch")
                    st.dataframe(result, width="stretch", hide_index=True)
                    st.download_button("تنزيل نتائج Clustering", to_csv_bytes(result), file_name="cluster_results.csv", mime="text/csv")
    if st.session_state.cluster_df is not None:
        st.caption("آخر نتائج التجميع محفوظة في الجلسة.")

with tabs[5]:
    st.markdown("### Anomaly Detection — اكتشاف الحالات غير المعتادة")
    st.write("يستخدم Isolation Forest على الأعمدة الرقمية لتحديد سجلات قد تكون غير معتادة إحصائيًا. هذه ليست بالضرورة أخطاء.")
    num_cols = active_df.select_dtypes(include=np.number).columns.tolist()
    if len(num_cols) < 1 or len(active_df) < 10:
        st.warning("تحتاج إلى عمود رقمي واحد على الأقل و10 صفوف.")
    else:
        anomaly_features = st.multiselect("اختر أعمدة الفحص", num_cols, default=num_cols[:min(5,len(num_cols))], key="anomaly_features")
        contamination = st.slider("النسبة المتوقعة للحالات الشاذة", 0.01, 0.20, 0.05, 0.01, key="contamination")
        if st.button("🔍 فحص الحالات الشاذة", width="stretch"):
            if not anomaly_features:
                st.error("اختر عمودًا رقميًا واحدًا على الأقل.")
            else:
                anomaly_source = active_df if len(active_df) <= 100_000 else active_df.sample(n=100_000, random_state=42)
                temp = anomaly_source[anomaly_features].replace([np.inf,-np.inf], np.nan).copy()
                temp = temp.fillna(temp.median(numeric_only=True))
                temp = temp.dropna()
                if len(temp) < 10:
                    st.error("الصفوف الصالحة بعد تجهيز البيانات أقل من 10.")
                else:
                    scaled = StandardScaler().fit_transform(temp)
                    detector = IsolationForest(contamination=contamination, random_state=42, n_estimators=200)
                    labels = detector.fit_predict(scaled)
                    out = anomaly_source.loc[temp.index].copy()
                    out["Anomaly_Status"] = np.where(labels == -1, "Anomaly", "Normal")
                    out["Anomaly_Score"] = detector.decision_function(scaled)
                    st.session_state.outlier_df = out.sort_values("Anomaly_Score")
                    n = int((labels == -1).sum())
                    a,b,c = st.columns(3)
                    with a: metric_card("السجلات المفحوصة", f"{len(temp):,}")
                    with b: metric_card("حالات غير معتادة", f"{n:,}")
                    with c: metric_card("النسبة", f"{n/len(temp)*100:.1f}%")
                    fig = px.histogram(out, x="Anomaly_Score", color="Anomaly_Status", title="Anomaly score distribution", template="plotly_dark")
                    fig.update_layout(paper_bgcolor="rgba(0,0,0,0)", plot_bgcolor="rgba(0,0,0,0)")
                    st.plotly_chart(fig, width="stretch")
                    st.dataframe(st.session_state.outlier_df.head(200), width="stretch", hide_index=True)
                    st.download_button("تنزيل نتائج الفحص", to_csv_bytes(st.session_state.outlier_df), file_name="anomaly_detection.csv", mime="text/csv")
    if st.session_state.outlier_df is not None:
        st.caption("آخر نتائج اكتشاف الشذوذ محفوظة في الجلسة.")

with tabs[6]:
    st.markdown("### التقارير والتصدير")
    st.write("جهّز كل ملف عند الطلب فقط لتقليل استهلاك الذاكرة أثناء رفع وتحليل البيانات الكبيرة.")
    report_df = active_df
    if st.button("📄 تجهيز تقرير جودة البيانات PDF", width="stretch"):
        try:
            with st.spinner("جاري تجهيز تقرير PDF..."):
                st.session_state.export_pdf_bytes = create_pdf_quality_report(report_df, st.session_state.source_name)
            st.success("تم تجهيز تقرير PDF.")
        except Exception as report_error:
            st.error(f"تعذر إنشاء تقرير PDF: {report_error}")
            st.session_state.export_pdf_bytes = make_report(report_df, st.session_state.source_name).encode("utf-8")
    if st.session_state.get("export_pdf_bytes"):
        st.download_button("⬇️ تنزيل تقرير الجودة", data=st.session_state.export_pdf_bytes,
                           file_name="AI_Data_Quality_Report.pdf" if st.session_state.get("export_pdf_bytes", b"").startswith(b"%PDF") else "data_quality_report.txt",
                           mime="application/pdf" if st.session_state.get("export_pdf_bytes", b"").startswith(b"%PDF") else "text/plain",
                           width="stretch")

    e1, e2, e3 = st.columns(3)
    with e1:
        if st.button("تجهيز CSV", width="stretch"):
            with st.spinner("جاري تجهيز CSV..."):
                st.session_state.export_csv_bytes = to_csv_bytes(report_df)
                gc.collect()
        if st.session_state.get("export_csv_bytes") is not None:
            st.download_button("تنزيل CSV", st.session_state.export_csv_bytes, file_name="dataset_export.csv", mime="text/csv", width="stretch")
    with e2:
        if st.button("📊 تجهيز Excel Dashboard", width="stretch"):
            try:
                with st.spinner("جاري إنشاء ملف Excel والرسوم..."):
                    st.session_state.export_excel_bytes = to_excel_bytes(report_df)
                    gc.collect()
            except Exception as export_error:
                st.error(f"تعذر إنشاء Excel: {export_error}. جرّب CSV إذا كان الملف قريبًا من حدود Excel.")
        if st.session_state.get("export_excel_bytes") is not None:
            st.download_button("تنزيل Excel مع Dashboard", st.session_state.export_excel_bytes, file_name="dataset_export.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", width="stretch")
    with e3:
        if st.button("تجهيز JSON", width="stretch"):
            with st.spinner("جاري تجهيز JSON..."):
                st.session_state.export_json_bytes = safe_frame(report_df).to_json(orient="records", force_ascii=False).encode("utf-8")
                gc.collect()
        if st.session_state.get("export_json_bytes") is not None:
            st.download_button("تنزيل JSON", st.session_state.export_json_bytes, file_name="dataset_export.json", mime="application/json", width="stretch")

    st.markdown("### إحصاءات وصفية")
    if len(report_df) > 200_000:
        st.caption("للحفاظ على الذاكرة، الإحصاءات الوصفية تعرض عينة عشوائية من 100,000 صف. عدد الصفوف الأصلي يظل كما هو.")
        stats_df = plot_sample(report_df, 100_000)
    else:
        stats_df = report_df
    st.dataframe(stats_df.describe(include="all").transpose(), width="stretch")

with tabs[7]:
    st.markdown("### SQL & Automation")
    st.write("صدّر نسخة من البيانات إلى قاعدة SQLite محلية، أو نزّل سكربت SQL لإنشاء جدول وإدخال البيانات.")
    sql_df = active_df
    safe_table = st.text_input("اسم الجدول", value="dataset_records")
    safe_table = "".join(ch if ch.isalnum() or ch == "_" else "_" for ch in safe_table)
    if not safe_table:
        safe_table = "dataset_records"
    if st.button("إنشاء قاعدة SQLite للتنزيل", width="stretch"):
        try:
            buffer = io.BytesIO()
            with sqlite3.connect(":memory:") as conn:
                sql_df.to_sql(safe_table, conn, index=False, if_exists="replace")
                backup = sqlite3.connect("file:memdb1?mode=memory&cache=shared", uri=True)
                conn.backup(backup)
                # Export using a temporary in-memory database file-like representation is not supported directly;
                # create a local temporary database file for the download.
                import tempfile
                with tempfile.NamedTemporaryFile(suffix=".db", delete=False) as tmp:
                    temp_path = tmp.name
                try:
                    with sqlite3.connect(temp_path) as disk_conn:
                        sql_df.to_sql(safe_table, disk_conn, index=False, if_exists="replace")
                    with open(temp_path, "rb") as f:
                        db_bytes = f.read()
                finally:
                    try: os.remove(temp_path)
                    except Exception: pass
                backup.close()
            st.download_button("⬇️ تنزيل قاعدة SQLite", db_bytes, file_name="dataset_database.sqlite", mime="application/vnd.sqlite3", width="stretch")
            st.success("تم تجهيز قاعدة SQLite.")
        except Exception as e:
            st.error(f"تعذر إنشاء SQLite: {e}")
    sql_table = safe_table.replace('"', '""')
    sql_lines = [f'CREATE TABLE "{sql_table}" (']
    cols = []
    for col in sql_df.columns:
        cname = str(col).replace('"', '""')
        dtype = "REAL" if pd.api.types.is_numeric_dtype(sql_df[col]) else "TEXT"
        cols.append(f'  "{cname}" {dtype}')
    sql_lines.append(",\n".join(cols))
    sql_lines.append(");")
    sql_script = "\n".join(sql_lines) + "\n\n-- Insert data from the exported CSV/Excel using your SQL client.\n"
    st.download_button("تنزيل SQL Schema", sql_script.encode("utf-8"), file_name="dataset_schema.sql", mime="application/sql", width="stretch")
    st.markdown("### سجل إجراءات الجلسة")
    activity = [
        {"Action": "Dataset loaded", "Status": "Done", "Details": st.session_state.source_name},
        {"Action": "Cleaned copy", "Status": "Ready" if st.session_state.clean_df is not None else "Not created", "Details": "Original data remains unchanged"},
        {"Action": "Machine Learning", "Status": "Trained" if st.session_state.model_info else "Not trained", "Details": str(st.session_state.model_info or "No model in this session")},
        {"Action": "Clustering", "Status": "Done" if st.session_state.cluster_df is not None else "Not run", "Details": "K-Means"},
        {"Action": "Anomaly detection", "Status": "Done" if st.session_state.outlier_df is not None else "Not run", "Details": "Isolation Forest"}
    ]
    st.dataframe(pd.DataFrame(activity), width="stretch", hide_index=True)

st.markdown("---")
st.markdown('<div class="small-muted">AI Data Intelligence Pro • النتائج استكشافية وتعتمد على جودة البيانات. راجع مخرجات النماذج قبل استخدامها في قرارات مهمة.</div>', unsafe_allow_html=True)
